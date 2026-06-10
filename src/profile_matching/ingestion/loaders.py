"""Filesystem document loaders.

Supports ``.txt``, ``.md``, ``.pdf`` and ``.docx``. PDF/DOCX support is optional
— if the relevant library is not installed, a clear, actionable error is raised
only when such a file is actually encountered (lazy import), keeping the core
package dependency-light.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from profile_matching.logging_config import get_logger
from profile_matching.utils.text import normalise_whitespace

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


@dataclass(frozen=True)
class LoadedDocument:
    """Raw text extracted from a source file."""

    path: Path
    text: str


def _load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - exercised only without dep
        raise RuntimeError(
            "Reading PDF resumes requires 'pypdf'. Install with: pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _load_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - exercised only without dep
        raise RuntimeError(
            "Reading DOCX resumes requires 'python-docx'. "
            "Install with: pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


_LOADERS = {
    ".txt": _load_txt,
    ".md": _load_txt,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
}


def load_document(path: str | Path) -> LoadedDocument:
    """Load and normalise a single document."""

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Document not found: {path}")

    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise ValueError(
            f"Unsupported file type '{suffix}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    text = normalise_whitespace(loader(path))
    logger.debug("Loaded %s (%d chars)", path.name, len(text))
    return LoadedDocument(path=path, text=text)


def load_documents(directory: str | Path) -> list[LoadedDocument]:
    """Load every supported document in ``directory`` (recursively).

    Files that fail to load are logged and skipped so a single corrupt file
    cannot abort an entire ingestion batch — important for production runs.
    """

    directory = Path(directory)
    if not directory.exists():
        raise FileNotFoundError(f"Directory not found: {directory}")

    documents: list[LoadedDocument] = []
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                documents.append(load_document(path))
            except Exception as exc:  # noqa: BLE001 - resilience over strictness
                logger.warning("Skipping %s: %s", path, exc)

    logger.info("Loaded %d documents from %s", len(documents), directory)
    return documents
